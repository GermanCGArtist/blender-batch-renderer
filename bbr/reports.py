"""Per-job data files in <job>/data/:

  frames.csv    one row per frame event, appended live (done or failed attempt)
  attempts.csv  one row per Blender run
  report.xlsx   summary + frames + attempts sheets with charts (rewritten at job end/stop)
"""
from __future__ import annotations

import csv
import os
import time
from typing import Optional

FRAME_COLS = ["frame", "status", "attempt_no", "attempts_for_frame", "finished_at", "render_s",
              "save_s", "blender_peak_mb", "process_peak_mb", "samples", "size_bytes", "file", "error"]
ATTEMPT_COLS = ["attempt", "started", "ended", "duration_s", "frames_rendered", "exit_code",
                "reason", "blender", "peak_rss_mb", "log"]


def _ts(t: Optional[float]) -> str:
    return time.strftime("%Y-%m-%d %H:%M:%S", time.localtime(t)) if t else ""


def _append(path, cols, row):
    new = not os.path.exists(path)
    with open(path, "a", newline="", encoding="utf-8") as f:
        w = csv.writer(f)
        if new:
            w.writerow(cols)
        w.writerow([row.get(c, "") for c in cols])


def append_frame(job, frame: int, rec: dict, attempt_no: int) -> None:
    try:
        job.data_dir.mkdir(parents=True, exist_ok=True)
        _append(job.data_dir / "frames.csv", FRAME_COLS, {
            "frame": frame, "status": rec.get("status"), "attempt_no": attempt_no,
            "attempts_for_frame": rec.get("attempts"),
            "finished_at": _ts(rec.get("finished_at") or time.time()),
            "render_s": _r(rec.get("render_s")), "save_s": _r(rec.get("save_s")),
            "blender_peak_mb": _r(rec.get("blender_peak_mb")),
            "process_peak_mb": _r(rec.get("process_peak_mb")),
            "samples": rec.get("samples") or "", "size_bytes": rec.get("size") or "",
            "file": rec.get("file") or "", "error": rec.get("last_error", "") if rec.get("status") != "done" else "",
        })
    except OSError:
        pass


def append_attempt(job, a: dict) -> None:
    try:
        job.data_dir.mkdir(parents=True, exist_ok=True)
        dur = (a["ended"] - a["started"]) if a.get("ended") and a.get("started") else None
        _append(job.data_dir / "attempts.csv", ATTEMPT_COLS, {
            "attempt": a.get("n"), "started": _ts(a.get("started")), "ended": _ts(a.get("ended")),
            "duration_s": _r(dur), "frames_rendered": a.get("frames_rendered"),
            "exit_code": a.get("exit_code"), "reason": a.get("reason"), "blender": a.get("blender"),
            "peak_rss_mb": _r(a.get("peak_rss_mb")), "log": a.get("log"),
        })
    except OSError:
        pass


def _r(v, nd=3):
    if v is None or v == "":
        return ""
    try:
        return round(float(v), nd)
    except (TypeError, ValueError):
        return v


def write_xlsx(job) -> Optional[str]:
    try:
        from openpyxl import Workbook
        from openpyxl.chart import LineChart, BarChart, Reference
        from openpyxl.styles import Font, PatternFill, Alignment
        from openpyxl.utils import get_column_letter
    except ImportError:
        return None
    from .parse import fmt_duration, fmt_bytes
    from .blendfile import version_str

    wb = Workbook()
    head_font = Font(bold=True, color="FFFFFF")
    head_fill = PatternFill("solid", fgColor="2F3B4C")

    def header(ws, cols):
        ws.append(cols)
        for c in ws[1]:
            c.font, c.fill = head_font, head_fill
            c.alignment = Alignment(vertical="center")
        ws.freeze_panes = "A2"

    def widths(ws, w):
        for i, x in enumerate(w, 1):
            ws.column_dimensions[get_column_letter(i)].width = x

    # ---------------- Frames ----------------
    wsf = wb.active
    wsf.title = "Frames"
    cols = ["Frame", "Status", "Attempts", "Render time (s)", "Save time (s)", "Blender peak mem (MB)",
            "Process peak mem (MB)", "Samples", "File size (MB)", "Finished at", "File / last error"]
    header(wsf, cols)
    recs = sorted(list(job.frames.items()), key=lambda kv: int(kv[0]))
    for k, r in recs:
        size = r.get("size")
        wsf.append([int(k), r.get("status"), r.get("attempts"), _r(r.get("render_s")), _r(r.get("save_s")),
                    _r(r.get("blender_peak_mb"), 1), _r(r.get("process_peak_mb"), 1), r.get("samples"),
                    _r(size / 1048576, 2) if size else None, _ts(r.get("finished_at")),
                    r.get("file") if r.get("status") == "done" else r.get("last_error", "")])
    widths(wsf, [8, 10, 9, 15, 13, 20, 20, 9, 13, 20, 60])
    n = len(recs)

    # ---------------- Attempts ----------------
    wsa = wb.create_sheet("Attempts")
    header(wsa, ["Attempt", "Started", "Ended", "Duration (s)", "Frames rendered", "Exit code",
                 "Result", "Blender", "Peak process mem (MB)", "Log file"])
    for a in job.attempts:
        dur = (a["ended"] - a["started"]) if a.get("ended") and a.get("started") else None
        wsa.append([a.get("n"), _ts(a.get("started")), _ts(a.get("ended")), _r(dur, 1),
                    a.get("frames_rendered"), a.get("exit_code"), a.get("reason"), a.get("blender"),
                    _r(a.get("peak_rss_mb"), 1), a.get("log")])
    widths(wsa, [9, 20, 20, 13, 15, 10, 50, 22, 20, 60])

    # ---------------- Summary ----------------
    ws = wb.create_sheet("Summary", 0)
    done = [r for _, r in recs if r.get("status") == "done"]
    times = [r["render_s"] for r in done if r.get("render_s")]
    si = job.scene_info or {}
    bu = job.blender_used or {}
    total_render = sum(times) if times else 0
    wall = None
    if job.attempts:
        wall = sum((a["ended"] - a["started"]) for a in job.attempts if a.get("ended") and a.get("started"))
    peaks = [r.get("process_peak_mb") for r in done if r.get("process_peak_mb")]
    bpeaks = [r.get("blender_peak_mb") for r in done if r.get("blender_peak_mb")]
    rows = [
        ("Job", job.name), ("Blend file", job.blend_path), ("Job folder", job.folder),
        ("Status", job.status + (f" – {job.status_note}" if job.status_note else "")),
        ("File saved with Blender", version_str(job.file_version)),
        ("Rendered with", f"{bu.get('label', '?')} ({bu.get('reason', '')})"),
        ("Scene", si.get("scene")), ("Engine", si.get("engine")), ("Device", si.get("device")),
        ("Resolution", f"{si.get('res_x')}×{si.get('res_y')} @ {si.get('res_pct')}%" if si else None),
        ("Samples", si.get("samples")), ("Frame range", f"{si.get('frame_start')}–{si.get('frame_end')} step {si.get('frame_step', 1)}" if si else None),
        ("FPS", round(si["fps"] / (si.get("fps_base") or 1), 3) if si.get("fps") else None),
        ("Output format", f"{si.get('file_format')} {si.get('color_depth') or ''}bit".strip() if si else None),
        ("Format note", si.get("format_note")),
        ("", ""),
        ("Frames total", job.total_frames), ("Frames done", len(done)),
        ("Frames failed", job.count("failed")), ("Blender runs (attempts)", len(job.attempts)),
        ("Total render time", fmt_duration(total_render) if times else None),
        ("Wall-clock time incl. loading", fmt_duration(wall) if wall else None),
        ("Average frame time (s)", round(total_render / len(times), 2) if times else None),
        ("Fastest frame (s)", round(min(times), 2) if times else None),
        ("Slowest frame (s)", round(max(times), 2) if times else None),
        ("Peak Blender memory (MB)", round(max(bpeaks), 1) if bpeaks else None),
        ("Peak process memory (MB)", round(max(peaks), 1) if peaks else None),
        ("Disk used by frames", fmt_bytes(job.used_bytes())),
        ("Preview video", job.preview_path),
        ("Report written", _ts(time.time())),
    ]
    for k, v in rows:
        ws.append([k, v])
        ws.cell(ws.max_row, 1).font = Font(bold=True)
        ws.cell(ws.max_row, 2).alignment = Alignment(horizontal="left")
    widths(ws, [30, 55])
    for sheet in (ws, wsf, wsa):
        sheet.sheet_properties.pageSetUpPr.fitToPage = True
        sheet.page_setup.fitToWidth = 1
        sheet.page_setup.fitToHeight = 0
        sheet.page_setup.orientation = "landscape"

    # ---------------- Charts ----------------
    if n >= 2:
        xref = Reference(wsf, min_col=1, min_row=2, max_row=n + 1)
        c1 = LineChart()
        c1.title, c1.y_axis.title, c1.x_axis.title = "Render time per frame", "seconds", "frame"
        c1.add_data(Reference(wsf, min_col=4, min_row=1, max_row=n + 1), titles_from_data=True)
        c1.set_categories(xref)
        c1.height, c1.width = 9, 22
        ws.add_chart(c1, "D2")
        c2 = LineChart()
        c2.title, c2.y_axis.title, c2.x_axis.title = "Memory per frame", "MB", "frame"
        c2.add_data(Reference(wsf, min_col=6, max_col=7, min_row=1, max_row=n + 1), titles_from_data=True)
        c2.set_categories(xref)
        c2.height, c2.width = 9, 22
        ws.add_chart(c2, "D21")
        c3 = BarChart()
        c3.title, c3.y_axis.title, c3.x_axis.title = "File size per frame", "MB", "frame"
        c3.add_data(Reference(wsf, min_col=9, min_row=1, max_row=n + 1), titles_from_data=True)
        c3.set_categories(xref)
        c3.height, c3.width = 9, 22
        ws.add_chart(c3, "D40")

    out = job.data_dir / "report.xlsx"
    tmp = job.data_dir / ".tmp_report.xlsx"
    job.data_dir.mkdir(parents=True, exist_ok=True)
    try:
        wb.save(tmp)
        os.replace(tmp, out)
    except PermissionError:
        # report.xlsx is open in Excel (Windows locks it) - write next to it instead
        alt = job.data_dir / f"report_{time.strftime('%Y%m%d_%H%M%S')}.xlsx"
        os.replace(tmp, alt)
        return str(alt)
    return str(out)
